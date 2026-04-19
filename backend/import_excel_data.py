#!/usr/bin/env python
"""
从桌面Excel导入全部物料/库位/库存数据
覆盖之前所有数据
"""
import os
import sys
import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'erp_backend.settings')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
django.setup()

import openpyxl
from decimal import Decimal
from datetime import datetime
from collections import defaultdict
from apps.inventory.models import Warehouse, Material, WarehouseLocation, Inventory

EXCEL_PATH = r'C:\Users\foxtion\Desktop\物料数据.xlsx'


def infer_category(name):
    """根据物料名称推断分类"""
    if not name:
        return '其他'
    name = str(name)
    if '模具' in name:
        return '硅胶模具'
    if '蜡烛' in name or '烛台' in name:
        return '蜡烛'
    if '杯垫' in name:
        return '杯垫'
    if '托盘' in name:
        return '托盘'
    if '收纳盒' in name or '收纳' in name:
        return '收纳盒'
    if '摆件' in name or '摆台' in name:
        return '摆件'
    if '吊坠' in name or '耳环' in name or '手镯' in name or '手环' in name:
        return '饰品'
    if '花瓶' in name:
        return '花瓶'
    if '梳子' in name:
        return '梳子'
    if '骰子' in name or '占卜' in name or '符文' in name:
        return '占卜'
    if '巧克力' in name:
        return '巧克力'
    if '书签' in name:
        return '书签'
    if '球' in name:
        return '球类'
    return '其他'


def parse_row(row):
    """解析Excel一行数据"""
    return {
        'product_code': row[0],
        'product_name': row[1],
        'location_code': row[2],
        'barcode': row[3],
        'outbound_barcode': row[4],
        'size': row[5],
        'aux': row[6],
        'is_empty': row[7] == '空位',
        'inbound_str': row[8],
        'qty': row[9],
        'out_status': row[10],
        'replenish': row[11],
        'auto_trigger': row[12],
        'reserve_pct': row[13],
        'manual_replenish': row[14],
        'manual_qty': row[15],
        'time_val': row[16],
        'datetime_val': row[17],
        'unit': row[18],
    }


def main():
    print('>>> 开始读取Excel文件...')
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    total_rows = ws.max_row - 1
    print(f'    总数据行数: {total_rows}')

    # 获取或创建默认仓库
    wh, _ = Warehouse.objects.get_or_create(
        code='MAIN',
        defaults={'name': '主仓库', 'location': '一楼', 'manager': '管理员', 'is_active': True}
    )
    print(f'    使用仓库: {wh.name}')

    # 清空旧数据
    print('>>> 清空旧数据...')
    Inventory.objects.all().delete()
    WarehouseLocation.objects.all().delete()
    Material.objects.all().delete()
    print('    Material/WarehouseLocation/Inventory 已清空')

    materials_map = {}  # code -> Material
    locations = []
    inventory_agg = defaultdict(lambda: {'qty': Decimal('0'), 'unit': '件', 'spec': ''})

    occupied_count = 0
    empty_count = 0

    print('>>> 解析并导入数据...')
    for idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True), 1):
        data = parse_row(row)
        loc_code = data['location_code']
        if not loc_code:
            continue

        # 处理size编码问题
        size = data['size'] or '小'
        if isinstance(size, str):
            # 处理可能的编码问题
            size = size.strip()
            if size not in ('小', '中', '大'):
                size = '小'

        # 处理入库时间
        inbound_time = None
        if data['datetime_val'] and isinstance(data['datetime_val'], datetime):
            inbound_time = data['datetime_val']

        # 处理数量
        qty = Decimal('0')
        if data['qty'] is not None:
            try:
                qty = Decimal(str(data['qty']))
            except:
                qty = Decimal('0')

        product_code = data['product_code'] or ''
        product_name = data['product_name'] or ''
        is_empty = data['is_empty']

        if not is_empty and product_code:
            occupied_count += 1

            # 收集物料（去重）
            if product_code not in materials_map:
                category = infer_category(product_name)
                mat = Material(
                    code=product_code,
                    name=product_name,
                    spec='',
                    category=category,
                    unit='件',
                    barcode=data['barcode'] or '',
                    qty=Decimal('0'),
                    status='active',
                    remark=''
                )
                materials_map[product_code] = mat

            # 汇总库存（按仓库+物料名称+规格）
            key = (wh.id, product_name, '')
            inventory_agg[key]['qty'] += qty
            inventory_agg[key]['unit'] = '件'
            inventory_agg[key]['spec'] = ''
        else:
            empty_count += 1

        # 构建库位对象
        locations.append(WarehouseLocation(
            warehouse=wh,
            location_code=str(loc_code).strip(),
            barcode=data['barcode'] or '',
            size=size,
            product_code=product_code,
            product_name=product_name,
            is_empty=is_empty,
            inbound_time=inbound_time,
            outbound_barcode=data['outbound_barcode'] or '',
            remark=''
        ))

        if idx % 1000 == 0:
            print(f'    已处理 {idx}/{total_rows} 行...')

    # 批量写入物料
    print(f'>>> 写入物料档案: {len(materials_map)} 条')
    Material.objects.bulk_create(list(materials_map.values()), batch_size=500)

    # 去重：同一库位保留有货的记录
    print(f'>>> 库位去重处理...')
    loc_dict = {}
    for loc in locations:
        key = loc.location_code
        if key not in loc_dict or (not loc.is_empty and loc_dict[key].is_empty):
            loc_dict[key] = loc
    locations = list(loc_dict.values())
    print(f'    去重后: {len(locations)} 条')

    # 批量写入库位
    empty_final = sum(1 for l in locations if l.is_empty)
    occupied_final = len(locations) - empty_final
    print(f'>>> 写入库位: {len(locations)} 条 (有货{occupied_final}, 空位{empty_final})')
    WarehouseLocation.objects.bulk_create(locations, batch_size=500)

    # 批量写入库存台账
    inventory_list = []
    for (wh_id, mat_name, spec), agg in inventory_agg.items():
        inventory_list.append(Inventory(
            warehouse_id=wh_id,
            material_name=mat_name,
            spec=spec,
            unit=agg['unit'],
            qty=agg['qty']
        ))
    print(f'>>> 写入库存台账: {len(inventory_list)} 条')
    Inventory.objects.bulk_create(inventory_list, batch_size=500)

    # 更新物料的qty为各仓库汇总
    print('>>> 同步物料库存数量...')
    for mat in Material.objects.all():
        total = Inventory.objects.filter(material_name=mat.name, spec=mat.spec or '').aggregate(s=__import__('django.db.models').db.models.Sum('qty'))['s'] or Decimal('0')
        mat.qty = total
        mat.save(update_fields=['qty'])

    # 同步库存预警状态
    print('>>> 同步库存预警状态...')
    for inv in Inventory.objects.all():
        inv.save()

    print('\n>>> 导入完成!')
    print(f'  物料档案: {Material.objects.count()}')
    print(f'  库位记录: {WarehouseLocation.objects.count()}')
    print(f'  库存台账: {Inventory.objects.count()}')
    print(f'  有货库位: {WarehouseLocation.objects.filter(is_empty=False).count()}')
    print(f'  空位库位: {WarehouseLocation.objects.filter(is_empty=True).count()}')
    print(f'  预警记录: {StockWarning.objects.count()}')


if __name__ == '__main__':
    main()
